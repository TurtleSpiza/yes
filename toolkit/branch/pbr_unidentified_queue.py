"""pbr_unidentified_queue.py - the identification queue: every register line whose contractor is still Unidentified (Tier 3),
grouped into supplier series, with ONE invoice to sight per series (the TechOne Document File that carries the attachment).

Reads the shipped register with python-calamine (rule 19.8); writes reports/Unidentified_Contractors_<ver>.md and .xlsx.
A series is (section, service, reference pattern): the same supplier numbers its invoices one way and bills one service,
so the pattern separates suppliers sharing a PK. The invoice to sight is the largest line in the series that carries a
TechOne attachment (Has Attachment = Y), else the largest line with the note that no attachment exists.
Usage: python3 pbr_unidentified_queue.py [register.xlsx] [outdir]
"""
import collections, datetime as dt, os, re, sys
from decimal import Decimal, ROUND_HALF_UP
from python_calamine import CalamineWorkbook
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
REG = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'registers', 'Parks_Branch_Transaction_Register_FY2627_v4.xlsx')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'reports')
D = lambda x: Decimal(str(x)).quantize(Decimal('0.01'), ROUND_HALF_UP)


def pattern(ref):
    """Reference shape: digits collapsed to #{n}, letters kept. 'INV-00043897' -> 'INV-#{8}', '00015355' -> '#{8}', '11913' -> '#{5}'."""
    s = str(ref).strip()
    return re.sub(r'\d+', lambda m: '#{%d}' % len(m.group(0)), s)


def main():
    wb = CalamineWorkbook.from_path(REG)
    cfg = {r[0]: r[1] for r in wb.get_sheet_by_name('Config').to_python(skip_empty_area=False) if r and r[0]}
    ver = str(cfg.get('WORKBOOK_VERSION', 'v?'))
    a, b = (int(x) for x in str(cfg['REGISTER_DATA']).split(':'))
    R = wb.get_sheet_by_name('Register').to_python(skip_empty_area=False)
    H = R[3]; ix = {h: i for i, h in enumerate(H)}
    g = lambda r, k: r[ix[k]]
    data = R[a - 1:b]
    total = sum(D(g(r, 'Amount ex GST')) for r in data)
    unid = [r for r in data if str(g(r, 'Contractor')).startswith('Unidentified')]
    series = collections.OrderedDict()
    for r in unid:
        key = (str(g(r, 'Section')), str(g(r, 'Service No')), pattern(g(r, 'Reference')))
        series.setdefault(key, []).append(r)

    def pick(rows):
        att = [r for r in rows if str(g(r, 'Src Has Attachment')).strip() == 'Y']
        pool = att or rows
        return max(pool, key=lambda r: abs(D(g(r, 'Amount ex GST')))), bool(att)

    def fmt_d(v):
        return v.strftime('%d-%b-%Y') if isinstance(v, (dt.date, dt.datetime)) else str(v)

    out = []
    for (sec, svc, pat), rows in series.items():
        amt = sum(D(g(r, 'Amount ex GST')) for r in rows)
        r0, has_att = pick(rows)
        refs = sorted({str(g(r, 'Reference')).strip() for r in rows})
        pks = collections.Counter(str(g(r, 'PK Charged')) for r in rows)
        nas = collections.Counter(str(g(r, 'Natural Account')) for r in rows)
        narrs = collections.Counter(str(g(r, 'Narration (GL line)') or '').split('\n')[0].strip()[:60] for r in rows)
        dates = sorted(g(r, 'Doc Date') for r in rows if isinstance(g(r, 'Doc Date'), (dt.date, dt.datetime)))
        dtype = collections.Counter(str(g(r, 'Doc Type')) for r in rows).most_common(1)[0][0]
        svc_name = ''
        out.append(dict(section=sec, service=svc, pattern=pat, lines=len(rows), amount=amt, refs=len(refs), first=fmt_d(dates[0]) if dates else '', last=fmt_d(dates[-1]) if dates else '',
                        pks='; '.join(f'{k} ({n})' for k, n in pks.most_common(4)), nas='; '.join(f'{k} ({n})' for k, n in nas.most_common(3)), doc_type=dtype,
                        narration=narrs.most_common(1)[0][0], route=str(g(r0, 'Procurement route (axis 4, v113)')),
                        sight=dict(docfile=str(g(r0, 'Src Document File')).replace('.0', ''), ref=str(g(r0, 'Reference')).strip(), date=fmt_d(g(r0, 'Doc Date')), amount=D(g(r0, 'Amount ex GST')),
                                   pk=str(g(r0, 'PK Charged')), na=str(g(r0, 'Natural Account')), attachment='Y' if has_att else 'N (no line in the series carries an attachment; request from AP)',
                                   linekey=str(g(r0, 'LineKey')), duid=str(g(r0, 'Src Document Unique ID')), narration=str(g(r0, 'Narration (GL line)') or '').replace('\n', ' | ')[:120],
                                   posted=fmt_d(g(r0, 'Src Posted Date')))))
    out.sort(key=lambda s: -abs(s['amount']))
    os.makedirs(OUT, exist_ok=True)
    stamp = dt.date.today().strftime('%d-%b-%Y')
    unid_amt = sum(s['amount'] for s in out)
    # ------------------------------------------------------------------ markdown
    md = [f'# Unidentified contractors, Parks Branch register FY2026/27 {ver} ({stamp})', '',
          f'**Position:** {len(unid):,} register lines carry no contractor evidence (Tier 3), ${unid_amt:,.2f} ex GST of the ${total:,.2f} register total, in {len(out)} supplier series.', '',
          'A series is one section, one service and one reference shape (the supplier\'s own invoice numbering), so two suppliers billing the same PK separate. '
          'The invoice to sight is the largest line in the series that carries a TechOne attachment: open the Document File in TechOne, sight the attachment, and the creditor name and ABN on it identify the whole series (rule 8 Tier 1 once the APLEDGER history is pulled for that creditor code).', '',
          '## Series by value', '',
          '| # | Section | Service | Ref shape | Lines | $ ex GST | Refs | Dates | PKs (lines) | Typical narration | Sight: Document File | Ref | Doc date | $ line | PK | Attachment |',
          '|---|---|---|---|---:|---:|---:|---|---|---|---|---|---|---:|---|---|']
    for k, s in enumerate(out, 1):
        t = s['sight']
        md.append(f"| {k} | {s['section']} | {s['service']} | `{s['pattern']}` | {s['lines']} | {s['amount']:,.2f} | {s['refs']} | {s['first']} to {s['last']} | {s['pks']} | {s['narration']} | **{t['docfile']}** | {t['ref']} | {t['date']} | {t['amount']:,.2f} | {t['pk']} | {t['attachment'][:1]} |")
    by_sec = collections.OrderedDict()
    for s in out:
        by_sec.setdefault(s['section'], [0, Decimal(0), 0]); by_sec[s['section']][0] += s['lines']; by_sec[s['section']][1] += s['amount']; by_sec[s['section']][2] += 1
    md += ['', '## By section', '', '| Section | Lines | $ ex GST | Series |', '|---|---:|---:|---:|']
    for sec, (n, amt, ns) in sorted(by_sec.items(), key=lambda kv: -abs(kv[1][1])):
        md.append(f'| {sec} | {n} | {amt:,.2f} | {ns} |')
    md += ['', '## How to use', '',
           '1. Work the table top down; the first ten series carry most of the value.',
           '2. For each series, open the Document File in TechOne (Financials > Documents), sight the attachment, and read the creditor name, ABN and creditor code.',
           '3. Request the APLEDGER creditor history for that code (one export per code, same criteria as Data_Acquisition F7 to F14) and add it to `toolkit/branch/pbr_histories_v4.json`; the next build identifies every line in the series (Method 12.0).',
           '4. Sight one invoice per series under rule 17 to confirm nature; the queue is regenerated from the register on every build.',
           '', f'Source: `{os.path.basename(REG)}`, Register rows {a}:{b}; generated by `toolkit/branch/pbr_unidentified_queue.py`.']
    open(os.path.join(OUT, f'Unidentified_Contractors_{ver}.md'), 'w', encoding='utf-8').write('\n'.join(md) + '\n')
    # ------------------------------------------------------------------ workbook
    xw = Workbook(); ws = xw.active; ws.title = 'Series'
    hdr_font = Font(name='Cambria', bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F3864'); MONEY = '$#,##0.00;($#,##0.00);"-"'
    cols = ['#', 'Section', 'Service', 'Reference shape', 'Lines', '$ ex GST', 'Distinct refs', 'First doc date', 'Last doc date', 'PKs (lines)', 'Natural accounts (lines)', 'Doc type', 'Typical narration', 'Procurement route',
            'SIGHT: Document File', 'SIGHT: Reference', 'SIGHT: Doc date', 'SIGHT: $ ex GST', 'SIGHT: PK', 'SIGHT: NA', 'SIGHT: Attachment in TechOne', 'SIGHT: LineKey', 'SIGHT: Document Unique ID', 'SIGHT: Narration', 'SIGHT: Posted']
    ws.append([f'Unidentified contractors, Parks Branch register FY2026/27 {ver} ({stamp}): {len(unid):,} lines, ${unid_amt:,.2f} ex GST, {len(out)} series. One invoice to sight per series.'])
    ws.append([]); ws.append(cols)
    for c in ws[3]:
        c.font = hdr_font; c.fill = fill
    for k, s in enumerate(out, 1):
        t = s['sight']
        ws.append([k, s['section'], s['service'], s['pattern'], s['lines'], float(s['amount']), s['refs'], s['first'], s['last'], s['pks'], s['nas'], s['doc_type'], s['narration'], s['route'],
                   t['docfile'], t['ref'], t['date'], float(t['amount']), t['pk'], t['na'], t['attachment'], t['linekey'], t['duid'], t['narration'], t['posted']])
        ws.cell(ws.max_row, 6).number_format = MONEY; ws.cell(ws.max_row, 18).number_format = MONEY
    for col, w in zip('ABCDEFGHIJKLMNOPQRSTUVWXY', [5, 18, 9, 16, 7, 14, 8, 12, 12, 34, 24, 18, 44, 28, 16, 14, 12, 14, 10, 8, 30, 30, 34, 60, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A4'
    wl = xw.create_sheet('Lines')
    lcols = ['Section', 'Service', 'Reference shape', 'LineKey', 'Reference', 'Doc Date', 'Doc Type', 'Amount ex GST', 'PK Charged', 'Natural Account', 'Narration (GL line)', 'Src Document File', 'Src Has Attachment', 'Src Document Unique ID', 'Contractor', 'Follow-up']
    wl.append(lcols)
    for c in wl[1]:
        c.font = hdr_font; c.fill = fill
    for (sec, svc, pat), rows in series.items():
        for r in rows:
            wl.append([sec, svc, pat, g(r, 'LineKey'), str(g(r, 'Reference')), fmt_d(g(r, 'Doc Date')), g(r, 'Doc Type'), float(D(g(r, 'Amount ex GST'))), g(r, 'PK Charged'), g(r, 'Natural Account'),
                       str(g(r, 'Narration (GL line)') or '').replace('\n', ' | '), str(g(r, 'Src Document File')).replace('.0', ''), g(r, 'Src Has Attachment'), g(r, 'Src Document Unique ID'), g(r, 'Contractor'), g(r, 'Follow-up')])
            wl.cell(wl.max_row, 8).number_format = MONEY
    wl.freeze_panes = 'A2'
    for col, w in zip('ABCDEFGHIJKLMNOP', [18, 9, 14, 30, 14, 12, 18, 14, 10, 8, 60, 14, 8, 34, 40, 60]):
        wl.column_dimensions[col].width = w
    xw.save(os.path.join(OUT, f'Unidentified_Contractors_{ver}.xlsx'))
    print(f'{len(unid)} lines, ${unid_amt:,.2f}, {len(out)} series -> {OUT}/Unidentified_Contractors_{ver}.md and .xlsx')
    for s in out[:12]:
        print(f"  {s['section']:18} {s['service']} {s['pattern']:14} {s['lines']:4} ${s['amount']:>12,.2f}  sight DocFile {s['sight']['docfile']} ref {s['sight']['ref']} ${s['sight']['amount']:,.2f} att {s['sight']['attachment'][:1]}")


if __name__ == '__main__':
    main()
