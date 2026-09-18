#!/usr/bin/env python3
"""
abn_bulk_verify.py

Bulk ABN verification against the Australian Business Register.

Reads a spreadsheet or CSV of supplier rows, runs the ATO checksum offline on
every ABN, then looks up each distinct ABN once on the ABR and writes a result
row per input row plus a summary.

Companion to the single-invoice `tax-invoice-compliance` skill. That skill
answers "is this one invoice payable". This answers "which of these several
hundred suppliers do I have a problem with".

Checks performed
    1. ATO checksum, offline, mathematical, no network
    2. ABN exists on the ABR
    3. ABN status is Active, and was Active as at the transaction date
    4. GST registration status, and whether registered as at the transaction date
    5. Supplier name on the record matches the ABR entity name, a registered
       business name, or a trading name

Lookup modes
    scrape  (default) parses the ABR public HTML page. Fine for a few hundred
            ABNs at a polite rate. No registration needed.
    api     uses the ABN Lookup web services JSON endpoint. Faster and
            supported for volume, but needs a free GUID from
            https://abr.business.gov.au/Tools/WebServices. Pass --guid.

Results are cached to JSON so a rerun costs nothing for ABNs already fetched.

Usage
    python abn_bulk_verify.py suppliers.xlsx --abn-col ABN
    python abn_bulk_verify.py suppliers.xlsx --abn-col ABN --name-col Supplier \
        --date-col "Invoice date" --out abn_results.xlsx
    python abn_bulk_verify.py ledger.csv --abn-col abn --guid <YOUR-GUID> --workers 8
    python abn_bulk_verify.py --abn 51824753556 33078502894      # ad hoc, no file

Exit codes
    0  every row cleared
    1  one or more rows returned a BLOCK verdict
    2  input or argument error
"""

from __future__ import annotations

import argparse
import concurrent.futures as cf
import difflib
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from datetime import datetime, date, timezone
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------- 1.0 constants

ABN_WEIGHTS = [10, 1, 3, 5, 7, 9, 11, 13, 15, 17, 19]

ABR_VIEW_URL = "https://abr.business.gov.au/ABN/View?abn={abn}"
ABR_JSON_URL = "https://abr.business.gov.au/json/AbnDetails.aspx?abn={abn}&guid={guid}"

HEADERS = {"User-Agent": "abn-bulk-verify/1.0 (+https://logan.qld.gov.au)"}

MONTHS = {m: i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
     "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], start=1)}

# Verdicts, ordered worst first. BLOCK verdicts stop payment.
BLOCK = {"CHECKSUM_FAIL", "NOT_FOUND", "ABN_CANCELLED", "ABN_NOT_ACTIVE_AT_DATE"}
FIXABLE = {"NOT_GST_REGISTERED", "GST_NOT_REGISTERED_AT_DATE", "NAME_MISMATCH"}
SEVERITY = ["CHECKSUM_FAIL", "NOT_FOUND", "ABN_CANCELLED", "ABN_NOT_ACTIVE_AT_DATE",
            "GST_NOT_REGISTERED_AT_DATE", "NOT_GST_REGISTERED", "NAME_MISMATCH",
            "LOOKUP_ERROR", "NO_ABN", "VALID"]


# ---------------------------------------------------------------- 2.0 checksum

def clean_abn(value: Any) -> str:
    """Pull 11 consecutive digits out of whatever the column contains."""
    if value is None:
        return ""
    s = re.sub(r"[^0-9]", "", str(value))
    if len(s) == 11:
        return s
    m = re.search(r"\d{11}", s)
    return m.group(0) if m else s


def validate_checksum(abn: str) -> bool:
    """ATO algorithm. Subtract 1 from the first digit, weight, sum, mod 89."""
    s = clean_abn(abn)
    if not s.isdigit() or len(s) != 11:
        return False
    digits = [int(c) for c in s]
    digits[0] -= 1
    return sum(d * w for d, w in zip(digits, ABN_WEIGHTS)) % 89 == 0


# ---------------------------------------------------------------- 3.0 fetching

def _get(url: str, timeout: int = 20, retries: int = 3, backoff: float = 1.8) -> str:
    last = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                raw = resp.read()
            try:
                return raw.decode("utf-8")
            except UnicodeDecodeError:
                return raw.decode("latin-1", errors="replace")
        except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, OSError) as e:
            last = e
            if attempt < retries - 1:
                time.sleep(backoff ** attempt)
    raise RuntimeError(f"fetch failed after {retries} attempts: {last}")


def _row_value(html: str, label: str) -> str | None:
    pattern = r"<th[^>]*>\s*" + re.escape(label) + r"\s*:?\s*</th>\s*<td[^>]*>(.*?)</td>"
    m = re.search(pattern, html, re.IGNORECASE | re.DOTALL)
    if not m:
        return None
    v = re.sub(r"<[^>]+>", " ", m.group(1))
    v = (v.replace("&nbsp;", " ").replace("&amp;", "&")
          .replace("&lt;", "<").replace("&gt;", ">"))
    return re.sub(r"\s+", " ", v).strip() or None


def _names_in(section: str) -> list[str]:
    out: list[str] = []
    for r in re.finditer(r"<tr[^>]*>\s*<td(?![^>]*colspan)[^>]*>(.*?)</td>\s*<td[^>]*>.*?</td>\s*</tr>",
                         section, re.DOTALL | re.IGNORECASE):
        t = re.sub(r"<[^>]+>", " ", r.group(1))
        t = t.replace("&nbsp;", " ").replace("&amp;", "&")
        t = re.sub(r"\s+", " ", t).strip()
        if not t or t in out or len(t) > 150:
            continue
        if re.match(r"^\d{1,2}\s+\w{3,9}\s+\d{4}$", t):
            continue
        out.append(t)
    return out


def lookup_scrape(abn: str) -> dict[str, Any]:
    url = ABR_VIEW_URL.format(abn=abn)
    rec: dict[str, Any] = {"abn": abn, "source": "abr-html", "source_url": url,
                           "fetched_at": datetime.now(timezone.utc).isoformat(),
                           "error": None}
    try:
        html = _get(url)
    except RuntimeError as e:
        rec["error"] = str(e)
        return rec

    rec["entity_name"] = _row_value(html, "Entity name")
    rec["abn_status_raw"] = _row_value(html, "ABN status")
    rec["entity_type"] = _row_value(html, "Entity type")
    rec["gst_raw"] = (_row_value(html, "Goods &amp; Services Tax (GST)")
                      or _row_value(html, "Goods & Services Tax (GST)"))
    rec["location"] = _row_value(html, "Main business location")

    bn = re.search(r"Business name\(s\)(.*?)(?:Trading name\(s\)|Deductible gift|ABN Lookup)",
                   html, re.IGNORECASE | re.DOTALL)
    tn = re.search(r"Trading name\(s\)(.*?)(?:Deductible gift|ABN Lookup|</body>|$)",
                   html, re.IGNORECASE | re.DOTALL)
    rec["business_names"] = _names_in(bn.group(1)) if bn else []
    rec["trading_names"] = _names_in(tn.group(1)) if tn else []
    if rec["entity_name"] is None and rec["abn_status_raw"] is None:
        rec["error"] = "not found on ABR"
    return rec


def lookup_api(abn: str, guid: str) -> dict[str, Any]:
    url = ABR_JSON_URL.format(abn=abn, guid=guid)
    rec: dict[str, Any] = {"abn": abn, "source": "abr-json", "source_url": url,
                           "fetched_at": datetime.now(timezone.utc).isoformat(),
                           "error": None}
    try:
        body = _get(url)
    except RuntimeError as e:
        rec["error"] = str(e)
        return rec
    m = re.search(r"\{.*\}", body, re.DOTALL)          # strip the JSONP wrapper
    if not m:
        rec["error"] = "unparseable JSONP response"
        return rec
    try:
        d = json.loads(m.group(0))
    except json.JSONDecodeError as e:
        rec["error"] = f"bad JSON: {e}"
        return rec
    if d.get("Message"):
        rec["error"] = d["Message"]
        return rec
    rec["entity_name"] = d.get("EntityName")
    rec["entity_type"] = d.get("EntityTypeName")
    status = d.get("AbnStatus") or ""
    eff = d.get("AbnStatusEffectiveFrom") or ""
    rec["abn_status_raw"] = f"{status} from {eff}".strip()
    rec["gst_raw"] = ("Registered from " + d["Gst"]) if d.get("Gst") else "Not currently registered for GST"
    rec["location"] = f"{d.get('AddressState','')} {d.get('AddressPostcode','')}".strip()
    rec["business_names"] = list(d.get("BusinessName") or [])
    rec["trading_names"] = []
    return rec


# ---------------------------------------------------------------- 4.0 parsing

def _parse_dmy(s: str) -> date | None:
    m = re.search(r"(\d{1,2})\s+([A-Za-z]{3})[a-z]*\s+(\d{4})", s)
    if m and m.group(2).title() in MONTHS:
        return date(int(m.group(3)), MONTHS[m.group(2).title()], int(m.group(1)))
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def parse_status(raw: str | None) -> tuple[str, date | None, date | None]:
    """Return (ACTIVE|CANCELLED|UNKNOWN, from_date, to_date)."""
    if not raw:
        return "UNKNOWN", None, None
    up = raw.upper()
    state = "ACTIVE" if "ACTIVE" in up else ("CANCELLED" if "CANCELLED" in up else "UNKNOWN")
    parts = re.split(r"\bto\b", raw, flags=re.IGNORECASE)
    return state, _parse_dmy(parts[0]), (_parse_dmy(parts[1]) if len(parts) > 1 else None)


def parse_gst(raw: str | None) -> tuple[bool, date | None, date | None]:
    """Return (currently_registered, from_date, to_date)."""
    if not raw:
        return False, None, None
    if "NOT CURRENTLY REGISTERED" in raw.upper():
        parts = re.split(r"\bto\b", raw, flags=re.IGNORECASE)
        return False, _parse_dmy(parts[0]), (_parse_dmy(parts[1]) if len(parts) > 1 else None)
    parts = re.split(r"\bto\b", raw, flags=re.IGNORECASE)
    frm = _parse_dmy(parts[0])
    to = _parse_dmy(parts[1]) if len(parts) > 1 else None
    return (to is None), frm, to


_SUFFIX = re.compile(
    r"\b(PTY|PTY\.|LTD|LTD\.|LIMITED|PROPRIETARY|INC|INCORPORATED|"
    r"THE TRUSTEE FOR|TRUSTEE FOR|ATF|T/A|TRADING AS|AUSTRALIA|AUST|GROUP|"
    r"HOLDINGS|SERVICES|CO|COMPANY|P/L)\b", re.IGNORECASE)


def normalise_name(s: str) -> str:
    s = (s or "").upper()
    s = re.sub(r"&", " AND ", s)
    s = _SUFFIX.sub(" ", s)
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def name_match(supplied: str, rec: dict[str, Any], threshold: float = 0.86) -> tuple[bool, float, str]:
    """Best match of the supplied name against every name on the ABR record."""
    target = normalise_name(supplied)
    if not target:
        return True, 1.0, ""                       # nothing to test against
    candidates = [rec.get("entity_name") or ""] + \
                 list(rec.get("business_names") or []) + \
                 list(rec.get("trading_names") or [])
    best, score, who = False, 0.0, ""
    for c in candidates:
        n = normalise_name(c)
        if not n:
            continue
        if n == target or n in target or target in n:
            return True, 1.0, c
        r = difflib.SequenceMatcher(None, target, n).ratio()
        if r > score:
            score, who = r, c
    return score >= threshold, round(score, 3), who


# ---------------------------------------------------------------- 5.0 verdict

def assess(abn: str, rec: dict[str, Any] | None, supplier: str | None,
           txn_date: date | None) -> dict[str, Any]:
    out: dict[str, Any] = {
        "ABN": abn or "", "Checksum": "", "Entity name": "", "Entity type": "",
        "ABN status": "", "GST status": "", "Location": "",
        "Name matched": "", "Match score": "", "Matched against": "",
        "Verdict": "", "Detail": "", "Source": "",
    }
    if not abn:
        out.update(Verdict="NO_ABN", Detail="No ABN supplied. Consider no-ABN withholding at 47%.")
        return out

    ok = validate_checksum(abn)
    out["Checksum"] = "Pass" if ok else "FAIL"
    if not ok:
        out.update(Verdict="CHECKSUM_FAIL",
                   Detail="Fails the ATO checksum, so the ABN cannot be valid.")
        return out

    if rec is None or rec.get("error"):
        err = (rec or {}).get("error", "no lookup performed")
        if "not found" in str(err).lower():
            out.update(Verdict="NOT_FOUND", Detail="Checksum passes but no ABR record exists.")
        else:
            out.update(Verdict="LOOKUP_ERROR", Detail=str(err))
        return out

    out["Entity name"] = rec.get("entity_name") or ""
    out["Entity type"] = rec.get("entity_type") or ""
    out["ABN status"] = rec.get("abn_status_raw") or ""
    out["GST status"] = rec.get("gst_raw") or ""
    out["Location"] = rec.get("location") or ""
    out["Source"] = rec.get("source_url") or ""

    state, a_from, a_to = parse_status(rec.get("abn_status_raw"))
    gst_now, g_from, g_to = parse_gst(rec.get("gst_raw"))

    verdicts: list[tuple[str, str]] = []
    if state == "CANCELLED":
        verdicts.append(("ABN_CANCELLED", f"ABN cancelled. {rec.get('abn_status_raw')}"))
    if txn_date:
        if a_from and txn_date < a_from:
            verdicts.append(("ABN_NOT_ACTIVE_AT_DATE",
                             f"ABN not active until {a_from:%d-%b-%Y}, transaction dated {txn_date:%d-%b-%Y}."))
        if a_to and txn_date > a_to:
            verdicts.append(("ABN_NOT_ACTIVE_AT_DATE",
                             f"ABN ceased {a_to:%d-%b-%Y}, transaction dated {txn_date:%d-%b-%Y}."))
        if g_from and txn_date < g_from:
            verdicts.append(("GST_NOT_REGISTERED_AT_DATE",
                             f"GST registration begins {g_from:%d-%b-%Y}, transaction dated {txn_date:%d-%b-%Y}."))
        if g_to and txn_date > g_to:
            verdicts.append(("GST_NOT_REGISTERED_AT_DATE",
                             f"GST registration ceased {g_to:%d-%b-%Y}, transaction dated {txn_date:%d-%b-%Y}."))
    if not gst_now and not any(v[0] == "GST_NOT_REGISTERED_AT_DATE" for v in verdicts):
        verdicts.append(("NOT_GST_REGISTERED",
                         "Not currently registered for GST. Any GST charged is not claimable."))

    if supplier:
        matched, score, who = name_match(supplier, rec)
        out["Name matched"] = "Yes" if matched else "NO"
        out["Match score"] = score
        out["Matched against"] = who
        if not matched:
            verdicts.append(("NAME_MISMATCH",
                             f"Supplied name does not match the ABR record. Closest, {who or 'none'}."))

    if not verdicts:
        out.update(Verdict="VALID", Detail="Active, GST registered, name reconciles.")
        return out
    verdicts.sort(key=lambda v: SEVERITY.index(v[0]))
    out.update(Verdict=verdicts[0][0], Detail="  ".join(v[1] for v in verdicts))
    return out


# ---------------------------------------------------------------- 6.0 cache

class Cache:
    def __init__(self, path: str | None, max_age_days: int):
        self.path = Path(path) if path else None
        self.max_age = max_age_days
        self.data: dict[str, Any] = {}
        if self.path and self.path.exists():
            try:
                self.data = json.loads(self.path.read_text())
            except json.JSONDecodeError:
                self.data = {}

    def get(self, abn: str) -> dict[str, Any] | None:
        rec = self.data.get(abn)
        if not rec:
            return None
        try:
            age = (datetime.now(timezone.utc) -
                   datetime.fromisoformat(rec["fetched_at"])).days
        except Exception:
            return None
        return rec if age <= self.max_age else None

    def put(self, abn: str, rec: dict[str, Any]) -> None:
        if not rec.get("error"):
            self.data[abn] = rec

    def save(self) -> None:
        if self.path:
            self.path.write_text(json.dumps(self.data, indent=1))


# ---------------------------------------------------------------- 7.0 io

def read_rows(path: str, sheet: str | None, header_row: int) -> tuple[list[dict], list[str]]:
    p = Path(path)
    if p.suffix.lower() in {".csv", ".txt", ".tsv"}:
        import csv
        delim = "\t" if p.suffix.lower() == ".tsv" else ","
        with p.open(newline="", encoding="utf-8-sig") as fh:
            rdr = csv.DictReader(fh, delimiter=delim)
            rows = list(rdr)
            return rows, list(rdr.fieldnames or [])
    try:
        import pandas as pd
    except ImportError:
        sys.exit("pandas is required to read spreadsheets. pip install pandas openpyxl")
    engine = "calamine"
    try:
        df = pd.read_excel(p, sheet_name=sheet or 0, engine=engine, header=header_row - 1)
    except Exception:
        df = pd.read_excel(p, sheet_name=sheet or 0, header=header_row - 1)
    df = df.dropna(how="all")
    return df.to_dict("records"), [str(c) for c in df.columns]


def coerce_date(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y", "%d %b %Y",
                "%Y-%m-%d %H:%M:%S", "%m/%d/%Y"):
        try:
            return datetime.strptime(s[:len(fmt) + 4], fmt).date()
        except ValueError:
            continue
    return _parse_dmy(s)


def write_xlsx(out_path: str, results: list[dict], summary: list[tuple[str, int]]) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    PANEL, NAVY, BAND = "1A2236", "1F3864", "F4F6FA"
    RED_L, AMBER_L, GREEN_L = "FFE0E0", "FFF2CC", "E2EFDA"
    FONT = "Segoe UI"

    wb = openpyxl.Workbook()

    sm = wb.active
    sm.title = "Summary"
    sm.sheet_view.showGridLines = False
    sm.cell(1, 1, "Bulk ABN verification").font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    for c in range(1, 5):
        sm.cell(1, c).fill = PatternFill("solid", fgColor=PANEL)
    sm.cell(2, 1, f"Run {datetime.now():%d-%b-%Y %H:%M}  ·  {len(results)} rows").font = \
        Font(name=FONT, size=9, italic=True, color="595959")
    for i, h in enumerate(["Verdict", "Rows", "Action"], 1):
        c = sm.cell(4, i, h); c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    ACTION = {
        "VALID": "No action.",
        "NAME_MISMATCH": "Confirm the trading name with the supplier, request a reissue if wrong.",
        "NOT_GST_REGISTERED": "Do not claim an input tax credit. Recover any GST charged.",
        "GST_NOT_REGISTERED_AT_DATE": "GST not claimable for that date. Query the supplier.",
        "ABN_CANCELLED": "Stop payment. Supplier cannot charge GST on a cancelled ABN.",
        "ABN_NOT_ACTIVE_AT_DATE": "Stop payment. Query the transaction date and the ABN.",
        "NOT_FOUND": "Stop payment. Checksum passes but the ABN is not on the register.",
        "CHECKSUM_FAIL": "Stop payment. The ABN is invalid on its face.",
        "NO_ABN": "Withhold at the top rate unless an exception applies.",
        "LOOKUP_ERROR": "Rerun. Network or register error, not a supplier fault.",
    }
    for i, (v, n) in enumerate(summary, 5):
        sm.cell(i, 1, v).font = Font(name=FONT, size=10, bold=v in BLOCK)
        sm.cell(i, 2, n).font = Font(name=FONT, size=10)
        sm.cell(i, 3, ACTION.get(v, "")).font = Font(name=FONT, size=10)
        shade = RED_L if v in BLOCK else (AMBER_L if v in FIXABLE else
                                          (GREEN_L if v == "VALID" else BAND))
        for c in range(1, 4):
            sm.cell(i, c).fill = PatternFill("solid", fgColor=shade)
    for col, w in zip("ABC", (30, 8, 70)):
        sm.column_dimensions[col].width = w

    rs = wb.create_sheet("Results")
    rs.sheet_view.showGridLines = False
    cols = list(results[0].keys()) if results else []
    for i, h in enumerate(cols, 1):
        c = rs.cell(1, i, h); c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name=FONT, size=9.5, bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    rs.row_dimensions[1].height = 24
    vi = cols.index("Verdict") + 1 if "Verdict" in cols else None
    for r, row in enumerate(results, 2):
        v = row.get("Verdict", "")
        shade = RED_L if v in BLOCK else (AMBER_L if v in FIXABLE else None)
        for i, k in enumerate(cols, 1):
            c = rs.cell(r, i, row.get(k))
            c.font = Font(name=FONT, size=9.5)
            if shade:
                c.fill = PatternFill("solid", fgColor=shade)
            elif r % 2 == 0:
                c.fill = PatternFill("solid", fgColor=BAND)
        if vi:
            rs.cell(r, vi).font = Font(name=FONT, size=9.5, bold=v in BLOCK)
        if "ABN" in cols:
            rs.cell(r, cols.index("ABN") + 1).number_format = "@"   # keep leading zeros
    widths = {"ABN": 14, "Checksum": 10, "Entity name": 38, "Entity type": 22,
              "ABN status": 24, "GST status": 26, "Location": 12,
              "Name matched": 12, "Match score": 11, "Matched against": 30,
              "Verdict": 26, "Detail": 62, "Source": 42}
    for i, k in enumerate(cols, 1):
        rs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(k, 18)
    rs.freeze_panes = "A2"
    rs.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(cols))}{len(results) + 1}"
    wb.save(out_path)


# ---------------------------------------------------------------- 8.0 main

def main() -> int:
    ap = argparse.ArgumentParser(description="Bulk ABN verification against the ABR.")
    ap.add_argument("input", nargs="?", help="xlsx, xlsm, csv or tsv file of supplier rows")
    ap.add_argument("--abn", nargs="*", default=None, help="verify these ABNs directly, no file")
    ap.add_argument("--abn-col", default="ABN")
    ap.add_argument("--name-col", default=None, help="supplier name column, enables the name match")
    ap.add_argument("--date-col", default=None, help="transaction date column, enables the as-at tests")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--out", default="abn_verification_results.xlsx")
    ap.add_argument("--guid", default=os.environ.get("ABR_GUID"),
                    help="ABN Lookup web services GUID, switches to the API mode")
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--delay", type=float, default=0.4, help="seconds between requests per worker")
    ap.add_argument("--cache", default=".abn_cache.json")
    ap.add_argument("--max-age-days", type=int, default=30)
    ap.add_argument("--threshold", type=float, default=0.86, help="name match ratio")
    ap.add_argument("--no-lookup", action="store_true", help="checksum only, no network")
    a = ap.parse_args()

    if not a.input and not a.abn:
        ap.error("supply an input file or --abn values")

    if a.abn:
        rows = [{a.abn_col: x} for x in a.abn]
        cols = [a.abn_col]
    else:
        rows, cols = read_rows(a.input, a.sheet, a.header_row)
        if a.abn_col not in cols:
            print(f"Column {a.abn_col!r} not found. Columns are: {cols}", file=sys.stderr)
            return 2

    for col, flag in ((a.name_col, "--name-col"), (a.date_col, "--date-col")):
        if col and col not in cols:
            print(f"Column {col!r} from {flag} not found. Columns are: {cols}", file=sys.stderr)
            return 2

    abns = sorted({clean_abn(r.get(a.abn_col)) for r in rows} - {""})
    print(f"{len(rows)} rows, {len(abns)} distinct ABNs")

    cache = Cache(a.cache, a.max_age_days)
    records: dict[str, dict[str, Any]] = {}
    to_fetch = []
    for x in abns:
        if not validate_checksum(x):
            continue
        hit = cache.get(x)
        if hit:
            records[x] = hit
        else:
            to_fetch.append(x)

    if to_fetch and not a.no_lookup:
        mode = "api" if a.guid else "scrape"
        print(f"fetching {len(to_fetch)} from the ABR, mode {mode}, {a.workers} workers")

        def job(x: str) -> tuple[str, dict[str, Any]]:
            time.sleep(a.delay)
            return x, (lookup_api(x, a.guid) if a.guid else lookup_scrape(x))

        done = 0
        with cf.ThreadPoolExecutor(max_workers=a.workers) as ex:
            for x, rec in ex.map(job, to_fetch):
                records[x] = rec
                cache.put(x, rec)
                done += 1
                if done % 25 == 0 or done == len(to_fetch):
                    print(f"  {done}/{len(to_fetch)}")
        cache.save()

    results = []
    for r in rows:
        abn = clean_abn(r.get(a.abn_col))
        supplier = str(r.get(a.name_col) or "") if a.name_col else None
        txn = coerce_date(r.get(a.date_col)) if a.date_col else None
        res = assess(abn, records.get(abn), supplier, txn)
        if a.name_col:
            res["Supplied name"] = supplier
        if a.date_col:
            res["Transaction date"] = txn.isoformat() if txn else ""
        results.append(res)

    counts: dict[str, int] = {}
    for res in results:
        counts[res["Verdict"]] = counts.get(res["Verdict"], 0) + 1
    summary = sorted(counts.items(), key=lambda kv: SEVERITY.index(kv[0]))

    print()
    for v, n in summary:
        tag = "BLOCK  " if v in BLOCK else ("REVIEW " if v in FIXABLE else "       ")
        print(f"  {tag}{v:28} {n}")

    write_xlsx(a.out, results, summary)
    print(f"\nwritten {a.out}")
    return 1 if any(v in BLOCK for v, _ in summary) else 0


if __name__ == "__main__":
    sys.exit(main())
